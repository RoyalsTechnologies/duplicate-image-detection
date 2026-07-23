"""Generate the DID Backend API group presentation PowerPoint deck."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUTPUT = Path(__file__).resolve().parents[1] / "DID_Backend_API_Presentation.pptx"

# Environmental / civic-tech palette
DARK = RGBColor(15, 61, 46)
ACCENT = RGBColor(34, 139, 104)
LIGHT = RGBColor(240, 248, 245)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(80, 100, 90)


def set_slide_bg(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(8.6), Inches(1.4))
    tf = title_box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.LEFT

    sub_box = slide.shapes.add_textbox(Inches(0.7), Inches(3.5), Inches(8.6), Inches(1.2))
    stf = sub_box.text_frame
    stf.text = subtitle
    sp = stf.paragraphs[0]
    sp.font.size = Pt(20)
    sp.font.color.rgb = LIGHT
    sp.alignment = PP_ALIGN.LEFT

    footer = slide.shapes.add_textbox(Inches(0.7), Inches(6.8), Inches(8.6), Inches(0.4))
    ftf = footer.text_frame
    ftf.text = "Alexander Adade · Royals Technologies"
    fp = ftf.paragraphs[0]
    fp.font.size = Pt(12)
    fp.font.color.rgb = ACCENT


def add_section_slide(prs: Presentation, title: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, ACCENT)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(8.4), Inches(1.2))
    tf = box.text_frame
    tf.text = title
    p = tf.paragraphs[0]
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER


def add_bullet_slide(
    prs: Presentation,
    title: str,
    bullets: list[str],
    *,
    subtitle: str | None = None,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    accent_bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.12))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = ACCENT
    accent_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(8.6), Inches(0.8))
    ttf = title_box.text_frame
    ttf.text = title
    tp = ttf.paragraphs[0]
    tp.font.size = Pt(30)
    tp.font.bold = True
    tp.font.color.rgb = DARK

    top = 1.35
    if subtitle:
        sub_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(8.6), Inches(0.5))
        stf = sub_box.text_frame
        stf.text = subtitle
        sp = stf.paragraphs[0]
        sp.font.size = Pt(14)
        sp.font.color.rgb = MUTED
        top = 1.75

    body_box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(8.4), Inches(5.0))
    btf = body_box.text_frame
    btf.word_wrap = True
    btf.vertical_anchor = MSO_ANCHOR.TOP

    for index, bullet in enumerate(bullets):
        paragraph = btf.paragraphs[0] if index == 0 else btf.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(18 if len(bullets) <= 6 else 16)
        paragraph.font.color.rgb = RGBColor(35, 45, 40)
        paragraph.space_after = Pt(10)


def add_two_column_slide(
    prs: Presentation,
    title: str,
    left_title: str,
    left_bullets: list[str],
    right_title: str,
    right_bullets: list[str],
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, WHITE)

    accent_bar = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(0.12))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = ACCENT
    accent_bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(8.6), Inches(0.8))
    ttf = title_box.text_frame
    ttf.text = title
    tp = ttf.paragraphs[0]
    tp.font.size = Pt(30)
    tp.font.bold = True
    tp.font.color.rgb = DARK

    for col_title, bullets, left_inch in (
        (left_title, left_bullets, 0.7),
        (right_title, right_bullets, 5.1),
    ):
        head = slide.shapes.add_textbox(Inches(left_inch), Inches(1.2), Inches(4.0), Inches(0.4))
        htf = head.text_frame
        htf.text = col_title
        hp = htf.paragraphs[0]
        hp.font.size = Pt(18)
        hp.font.bold = True
        hp.font.color.rgb = ACCENT

        body = slide.shapes.add_textbox(Inches(left_inch), Inches(1.65), Inches(4.0), Inches(4.8))
        btf = body.text_frame
        btf.word_wrap = True
        for index, bullet in enumerate(bullets):
            paragraph = btf.paragraphs[0] if index == 0 else btf.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(15)
            paragraph.font.color.rgb = RGBColor(35, 45, 40)
            paragraph.space_after = Pt(8)


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "Duplicate Image Detection (DID)",
        "Backend API for environmental and public-concern reporting",
    )

    add_bullet_slide(
        prs,
        "Problem Statement",
        [
            "Citizens report flooding, refuse dumps, potholes, pollution, and sanitation issues.",
            "The same problem is often reported many times with similar photos.",
            "Duplicate reports waste response-team time and clutter the dataset.",
            "Manual review does not scale as report volume grows.",
        ],
        subtitle="Why we built this system",
    )

    add_bullet_slide(
        prs,
        "Our Solution",
        [
            "Accept photo + GPS location + category submissions through a web UI and REST API.",
            "Detect likely duplicate reports using image, location, and category signals.",
            "Queue uncertain matches for admin review instead of auto-rejecting them.",
            "Validate uploads with computer vision so irrelevant images are rejected.",
            "Auto-suggest descriptions to make reporting faster for citizens.",
        ],
        subtitle="DID Backend API",
    )

    add_section_slide(prs, "How It Works")

    add_bullet_slide(
        prs,
        "User Flow",
        [
            "1. Citizen opens the report page and selects a photo of the issue.",
            "2. The system suggests a description (user can edit it).",
            "3. User pins location and chooses a category.",
            "4. Report is submitted to POST /api/v1/reports.",
            "5. API stores the image, analyzes it, and runs duplicate detection.",
            "6. Result: new report, duplicate, or possible duplicate for review.",
        ],
    )

    add_bullet_slide(
        prs,
        "Duplicate Detection Pipeline",
        [
            "Step 1: Extract SHA-256 hash, perceptual hash, and image embedding.",
            "Step 2: Search nearby reports in PostgreSQL + PostGIS within a radius.",
            "Step 3: Compare vector similarity (pgvector) and perceptual hash scores.",
            "Step 4: Apply category and time-window rules.",
            "Outcomes: new · duplicate · possible_duplicate · supporting_evidence",
        ],
        subtitle="Multi-signal matching",
    )

    add_two_column_slide(
        prs,
        "Duplicate Signals",
        "Image signals",
        [
            "SHA-256 exact file match",
            "Perceptual hash similarity",
            "CLIP / YOLO embedding vectors",
            "Detected objects in scene",
        ],
        "Context signals",
        [
            "GPS distance (PostGIS)",
            "Report category match",
            "Configurable time window",
            "Confidence score + evidence payload",
        ],
    )

    add_section_slide(prs, "Architecture")

    add_bullet_slide(
        prs,
        "System Architecture",
        [
            "Web UI (home page) → FastAPI API on port 3050",
            "PostgreSQL + PostGIS + pgvector on host port 3060",
            "Redis on host port 3061 for rate limits and hash cache",
            "Computer vision providers: local, embedding, yolov11, yolov11-cls",
            "Vision narration service for auto-generated descriptions",
            "Docker Compose for local development and production deployment",
        ],
    )

    add_two_column_slide(
        prs,
        "Tech Stack",
        "Core platform",
        [
            "FastAPI + Uvicorn",
            "PostgreSQL + PostGIS",
            "pgvector",
            "Redis",
            "Alembic migrations",
            "Docker / Docker Compose",
        ],
        "Intelligence layer",
        [
            "Local CV heuristics",
            "CLIP embeddings",
            "YOLOv11 detection",
            "YOLOv11 classification",
            "Groq vision narration",
            "ML training pipeline (ml/)",
        ],
    )

    add_bullet_slide(
        prs,
        "Report Categories",
        [
            "refuse_dump — illegal dumping and litter",
            "blocked_drain — blocked gutters and drains",
            "flooding — standing water and flood damage",
            "pothole — road surface damage",
            "pollution — smoke, air, and environmental pollution",
            "broken_public_facility — damaged public infrastructure",
            "sanitation — dirty or unsanitary public areas",
            "other — anything not covered above",
        ],
    )

    add_section_slide(prs, "Computer Vision")

    add_bullet_slide(
        prs,
        "Computer Vision Providers",
        [
            "local — fast heuristics, no heavy ML dependencies (default)",
            "embedding — CLIP semantic embeddings + relevance scoring",
            "yolov11 — object detection for scene understanding",
            "yolov11-cls — image classification per concern type",
            "Irrelevant images are rejected before a report is accepted.",
            "Fine-tuned weights can be trained via python -m ml.pipeline.",
        ],
    )

    add_section_slide(prs, "Live Demo")

    add_bullet_slide(
        prs,
        "Demo Script",
        [
            "Open http://localhost:3050/ (or production URL).",
            "Upload a photo of flooding, refuse, or pollution.",
            "Show the auto-generated description appearing in the form.",
            "Set latitude/longitude and choose a matching category.",
            "Submit the report and show the success result panel.",
            "Submit the same photo again nearby → duplicate detected.",
            "Optional: open /docs for the OpenAPI technical view.",
        ],
        subtitle="5-minute walkthrough",
    )

    add_two_column_slide(
        prs,
        "Key API Endpoints",
        "Public / user",
        [
            "GET / — report upload UI",
            "GET /health — health check",
            "POST /api/v1/describe-image",
            "POST /api/v1/reports",
            "GET /api/v1/reports/{id}",
        ],
        "Admin / review",
        [
            "GET /api/v1/duplicate-reviews",
            "POST /api/v1/duplicate-reviews/{id}/resolve",
            "PATCH /api/v1/reports/{id}/status",
            "GET /api/v1/reports/{id}/duplicates",
            "IP whitelist + rate limiting on protected routes",
        ],
    )

    add_bullet_slide(
        prs,
        "Security & Reliability",
        [
            "IP whitelist on all routes except / and /health.",
            "Upload rate limiting backed by Redis.",
            "Maximum upload size enforced per request.",
            "CV relevance checks reject non-concern images.",
            "Database migrations run automatically on container startup.",
            "Production uses custom ports to avoid conflicts with existing services.",
        ],
    )

    add_two_column_slide(
        prs,
        "Challenges & Solutions",
        "Challenge",
        [
            "Port 5432 / 6379 already used on production",
            "Missing ML weights broke Docker startup",
            "Slow first CV request after deploy",
            "Large photos exceeded narration API limits",
        ],
        "Solution",
        [
            "Mapped Postgres to 3060 and Redis to 3061",
            "Bake stock YOLO weights into runtime-cv image",
            "Background CV warmup on application startup",
            "Compress images before vision API calls",
        ],
    )

    add_bullet_slide(
        prs,
        "Results & Impact",
        [
            "Citizens can report issues faster with auto-suggested descriptions.",
            "Duplicate reports are flagged before they flood the workflow.",
            "Uncertain matches go to a review queue instead of being lost.",
            "Image validation improves overall data quality.",
            "Dockerized deployment makes the system portable to production.",
        ],
    )

    add_bullet_slide(
        prs,
        "Next Steps",
        [
            "Admin dashboard for duplicate review and report management.",
            "Mobile app integration via the REST API.",
            "S3 object storage for uploaded images.",
            "Fine-tuned YOLO models for local concern types.",
            "Analytics dashboard for hotspot mapping and trend reporting.",
        ],
        subtitle="Future roadmap",
    )

    add_title_slide(
        prs,
        "Thank You",
        "Questions & discussion",
    )

    return prs


def main() -> None:
    prs = build_presentation()
    prs.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
